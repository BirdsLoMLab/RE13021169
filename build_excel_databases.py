#!/usr/bin/env python3
"""Export the LOM master JSON databases to full Excel workbooks.

Produces three workbooks in battlesim/reference/excel/:
  LOM_Database_Master.xlsx        - one sheet per entity database (all scalar fields)
  LOM_Database_Progressions.xlsx  - every level/star/skin scaling table, one row per level
  LOM_Database_Bosses.xlsx        - all PvE boss stat tables
"""
import json, os, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(ROOT, "battlesim", "reference")
OUT = os.path.join(REF, "excel")
os.makedirs(OUT, exist_ok=True)
CELL_MAX = 32000
TODAY = datetime.date.today().isoformat()

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")


def load(name):
    with open(os.path.join(REF, name), encoding="utf-8") as fh:
        return json.load(fh)


# ---------------------------------------------------------------- attribute map
def attribute_map():
    """ID -> (key, description) parsed from 03_ATTRIBUTES.md markdown tables."""
    out = {}
    with open(os.path.join(REF, "03_ATTRIBUTES.md"), encoding="utf-8") as fh:
        for line in fh:
            if not line.startswith("|"):
                continue
            cells = [c.strip().strip("*") for c in line.strip().strip("|").split("|")]
            if len(cells) < 2 or not re.fullmatch(r"\d{3,5}", cells[0]):
                continue
            key = cells[1] if cells[1] not in ("—", "-", "") else ""
            desc = cells[-1]
            if cells[0] not in out or (key and not out[cells[0]][0]):
                out[cells[0]] = (key, desc)
    return out


ATTRS = attribute_map()


def attr_label(aid):
    key, desc = ATTRS.get(str(aid), ("", ""))
    return key or desc or ""


def pairs(value):
    """Render [[id, val], ...] as 'key(id)=val; ...'; fall back to JSON."""
    if value is None:
        return ""
    if isinstance(value, list) and value and all(
        isinstance(p, (list, tuple)) and len(p) == 2 and isinstance(p[0], int) for p in value
    ):
        parts = []
        for aid, val in value:
            lbl = attr_label(aid)
            parts.append(f"{lbl}({aid})={val}" if lbl else f"{aid}={val}")
        return "; ".join(parts)
    return cell(value)


def idpairs(value):
    """Render [[id, value], ...] as 'id=value' without attribute-name decoding."""
    if value is None:
        return ""
    if isinstance(value, list) and value and all(
        isinstance(p, (list, tuple)) and len(p) == 2 and not isinstance(p[0], (list, dict))
        and not isinstance(p[1], (list, dict)) for p in value
    ):
        return "; ".join(f"{a}={b}" for a, b in value)
    return cell(value)


def cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float, str)):
        s = value if not isinstance(value, str) else value
        if isinstance(s, str) and len(s) > CELL_MAX:
            return s[:CELL_MAX] + " …[truncated]"
        return s
    s = json.dumps(value, ensure_ascii=False)
    return s[:CELL_MAX] + " …[truncated]" if len(s) > CELL_MAX else s


def sheet(wb, title, headers, rows, widths=None, note=None, keep_empty=False):
    if not keep_empty and rows:
        keep = [i for i in range(len(headers))
                if any(i < len(r) and r[i] not in ("", None) for r in rows)]
        if len(keep) != len(headers):
            headers = [headers[i] for i in keep]
            rows = [[r[i] if i < len(r) else "" for i in keep] for r in rows]
    ws = wb.create_sheet(title[:31])
    r = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, size=9, color="555555")
        r = 2
    for c, h in enumerate(headers, 1):
        x = ws.cell(r, c, h)
        x.fill, x.font = HDR_FILL, HDR_FONT
        x.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    head = (2 if note else 1)
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(headers))}{max(head, r)}"
    for c, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            sample = max([len(str(h))] + [len(str(row[c - 1])) for row in rows[:200] if c <= len(row)] or [10])
            w = min(max(sample + 2, 10), 60)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[head].height = 28
    return ws


def readme(wb, title, lines):
    ws = wb.create_sheet(title)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    r = 1
    for a, b in lines:
        ws.cell(r, 1, a).font = Font(bold=True) if a and not b else Font()
        if a and b:
            ws.cell(r, 1, a).font = Font(bold=True)
        ws.cell(r, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return ws


def flat_rows(records, columns, renderers=None):
    renderers = renderers or {}
    rows = []
    for rec in records:
        rows.append([renderers.get(c, cell)(rec.get(c)) for c in columns])
    return rows


# ================================================================= load sources
pals = load("pals_master.json")
skills = load("skills_master.json")
mounts = load("mounts_master.json")
heroes = load("star_heroes_master.json")
avians = load("avians_master.json")
affixes = load("avian_affixes_master.json")
artifacts = load("artifacts_master.json")
backs = load("back_accessories_master.json")
bosses = load("bosses_master.json")

# ================================================================= 1. MASTER
wb = Workbook()
wb.remove(wb.active)

readme(wb, "README", [
    ("Legend of Mushroom — Master Database", ""),
    ("Generated", TODAY),
    ("Source", "battlesim/reference/*_master.json (LOM_Database-5.xlsx + decoded game config bundle)"),
    ("Decoder", "decode_config_data.py — 77-field unit configs, XOR config key 24455"),
    ("Credit", "Bird → Discord @birrrd08"),
    ("", ""),
    ("Sheets", ""),
    ("Pals", f"{pals['_total']} pals — IDs, race, damage multiplier, skill pairs"),
    ("Skills", f"{skills['_total']} active skills — IDs, cooldown, targeting, effect IDs"),
    ("Mounts", f"{mounts['_total_xlsx_mounts']} mounts ({mounts['_combat_skill_mounts']} with combat skills, {mounts['_tier_mounts_no_skill']} tier mounts)"),
    ("Mounts (config only)", "Mounts present in game config but not in the player-facing database"),
    ("Star Heroes", f"{heroes['_total']} star heroes (ConfigAngel) — {heroes['implemented']} implemented in the sim"),
    ("Star Heroes (config only)", "Angels present in config but not yet in the database"),
    ("Avians", f"{avians['_total']} avians (ConfigFly) — there is no ConfigBird"),
    ("Avian Affixes", f"{affixes['_total']} breeding affixes/entries"),
    ("Artifacts", f"{artifacts['_total']} artifact skins"),
    ("Back Accessories", f"{backs['_total']} back decorations"),
    ("Attribute IDs", "Attribute ID registry used by every attr/value_plus column"),
    ("Constants", "Battle, PvP and attribute-cap constants from data/constants/*.json"),
    ("", ""),
    ("Unit conventions", "num_type 2 attributes are /10000 (8000 = 80%). att_dam / skill_dam_extra are /10000 multipliers (100000 = 10x). att_speed /10000 (14000 = 1.4 attacks per cycle)."),
    ("Nested data", "Level-by-level scaling lives in LOM_Database_Progressions.xlsx. Boss stat tables live in LOM_Database_Bosses.xlsx."),
    ("Known quirks", "Inspire/Suppress names are swapped in the game code and preserved here. Guardian Spirits (ConfigSpirit) and Star Heroes (ConfigAngel) are separate systems."),
])

sheet(wb, "Pals",
      ["index", "name", "rarity", "family_type", "race_id", "race_name", "pet_id", "unit_id",
       "quality", "dmg_multiplier", "dmg_multiplier_raw", "power", "skill_pairs", "if_activity",
       "if_season", "open_time", "effect"],
      flat_rows(pals["pals"],
                ["index", "name", "rarity", "family_type", "race_id", "race_name", "pet_id", "unit_id",
                 "quality", "dmg_multiplier", "dmg_multiplier_raw", "power", "skill_pairs", "if_activity",
                 "if_season", "open_time", "effect"]),
      widths={"effect": 80, "name": 24})

sheet(wb, "Skills",
      ["index", "name", "rarity", "skill_id", "skill_type", "quality", "cooldown", "release_time",
       "release_interval", "target_type", "target_range", "skill_effect_ids", "buff_group_ids",
       "initial_power", "max_power", "power_recovery", "event_exclusive", "if_activity", "if_season",
       "effect", "passive_effect"],
      flat_rows(skills["skills"],
                ["index", "name", "rarity", "skill_id", "skill_type", "quality", "cooldown", "release_time",
                 "release_interval", "target_type", "target_range", "skill_effect_ids", "buff_group_ids",
                 "initial_power", "max_power", "power_recovery", "event_exclusive", "if_activity", "if_season",
                 "effect", "passive_effect"]),
      widths={"effect": 80, "passive_effect": 40, "name": 24},
      note="initial_power / max_power / power_recovery are XOR-encoded raw config values (see *_note fields in the JSON source).")

mrows = []
for m in mounts["mounts"]:
    ap = m.get("ability_progression") or {}
    mrows.append([cell(m.get(k)) for k in
                  ["index", "name", "rarity", "mount_id", "quality_code", "mount_type", "is_tier_mount",
                   "min_speed", "max_speed", "power", "skill_id", "skill_name", "max_skin_level",
                   "if_activity", "open_time", "effect", "passive_effect"]]
                 + [ap.get("ability_id", ""), ap.get("total_levels", "")])
sheet(wb, "Mounts",
      ["index", "name", "rarity", "mount_id", "quality_code", "mount_type", "is_tier_mount",
       "min_speed", "max_speed", "power", "skill_id", "skill_name", "max_skin_level",
       "if_activity", "open_time", "effect", "passive_effect", "ability_id", "ability_levels"],
      mrows, widths={"effect": 80, "passive_effect": 36, "name": 24})

sheet(wb, "Mounts (config only)",
      ["mount_id", "name", "quality_code", "mount_type", "min_speed", "max_speed", "power",
       "if_activity", "open_time"],
      flat_rows(mounts["config_only_mounts"],
                ["mount_id", "name", "quality_code", "mount_type", "min_speed", "max_speed", "power",
                 "if_activity", "open_time"]))

sheet(wb, "Star Heroes",
      ["index", "name", "rarity", "angel_id", "angel_type", "quality", "max_stars", "cost",
       "battle_skill1_cost", "battle_skill2_cost", "main_type", "support_type", "skill_ids",
       "limited", "implemented", "open_time", "main_skill", "support_skill"],
      flat_rows(heroes["star_heroes"],
                ["index", "name", "rarity", "angel_id", "angel_type", "quality", "max_stars", "cost",
                 "battle_skill1_cost", "battle_skill2_cost", "main_type", "support_type", "skill_ids",
                 "limited", "implemented", "open_time", "main_skill", "support_skill"]),
      widths={"main_skill": 70, "support_skill": 70, "name": 24})

sheet(wb, "Star Heroes (config only)",
      ["angel_id", "name", "type", "quality", "max_stars", "skill_ids", "time_limit", "open_time"],
      flat_rows(heroes["config_only_angels"],
                ["angel_id", "name", "type", "quality", "max_stars", "skill_ids", "time_limit", "open_time"]))

sheet(wb, "Avians",
      ["index", "name", "config_name", "rarity", "fly_id", "unit_id", "skill_id", "quality", "fly_type",
       "hybrid_type", "collab_flag", "attr", "event_exclusive", "collaboration", "limited_breeding",
       "food_required", "fly_special", "open_time", "effect", "home_weight", "entry_weight"],
      flat_rows(avians["avians"],
                ["index", "name", "config_name", "rarity", "fly_id", "unit_id", "skill_id", "quality", "fly_type",
                 "hybrid_type", "collab_flag", "attr", "event_exclusive", "collaboration", "limited_breeding",
                 "food_required", "fly_special", "open_time", "effect", "home_weight", "entry_weight"],
                {"attr": pairs, "home_weight": idpairs, "entry_weight": idpairs}),
      widths={"effect": 70, "attr": 44, "home_weight": 50, "entry_weight": 50, "name": 24},
      note="home_weight / entry_weight are breeding weight tables: id=weight pairs (0 = cannot roll).")

sheet(wb, "Avian Affixes",
      ["id", "name", "rarity", "entry_id", "entry_rarity_code", "max_level", "belong_talent",
       "power_at_max_level", "passive_skill_ids", "special_effect_ids", "home_effect", "conflict_entry",
       "desc_parm", "description"],
      flat_rows(affixes["avian_affixes"],
                ["id", "name", "rarity", "entry_id", "entry_rarity_code", "max_level", "belong_talent",
                 "power_at_max_level", "passive_skill_ids", "special_effect_ids", "home_effect", "conflict_entry",
                 "desc_parm", "description"]),
      widths={"description": 60, "name": 26})

sheet(wb, "Artifacts",
      ["index", "name", "rarity", "artifact_id", "quality_code", "skin_skill_id", "skin_skill_name",
       "skin_skill_ids_all_levels", "if_activity", "open_time", "effect", "passive_effect"],
      flat_rows(artifacts["artifacts"],
                ["index", "name", "rarity", "artifact_id", "quality_code", "skin_skill_id", "skin_skill_name",
                 "skin_skill_ids_all_levels", "if_activity", "open_time", "effect", "passive_effect"]),
      widths={"effect": 90, "passive_effect": 30, "name": 26})

sheet(wb, "Back Accessories",
      ["index", "name", "rarity", "back_id", "skill_id", "effect", "passive_effect"],
      flat_rows(backs["back_accessories"],
                ["index", "name", "rarity", "back_id", "skill_id", "effect", "passive_effect"]),
      widths={"effect": 90, "passive_effect": 30, "name": 26})

sheet(wb, "Attribute IDs", ["id", "key", "description"],
      [[int(k), v[0], v[1]] for k, v in sorted(ATTRS.items(), key=lambda kv: int(kv[0]))],
      widths={"description": 80})

# constants
crows = []
for fname in ("battle_constants", "pvp_constants", "attribute_caps"):
    with open(os.path.join(ROOT, "data", "constants", fname + ".json"), encoding="utf-8") as fh:
        data = json.load(fh)
    notes = data.get("_notes") or {}
    for k, v in data.items():
        if k == "_notes":
            continue
        crows.append([fname, k, cell(v), cell(notes.get(k, ""))])
sheet(wb, "Constants", ["source", "key", "value", "note"], crows,
      widths={"value": 60, "note": 70, "key": 34},
      note="Battle / PvP / attribute-cap constants decoded from the game config (data/constants/*.json).")

wb.save(os.path.join(OUT, "LOM_Database_Master.xlsx"))
print("wrote LOM_Database_Master.xlsx")

# ================================================================= 2. PROGRESSIONS
wb = Workbook()
wb.remove(wb.active)
readme(wb, "README", [
    ("Legend of Mushroom — Progression Tables", ""),
    ("Generated", TODAY),
    ("Source", "battlesim/reference/*_master.json — every level / star / skin scaling array, one row per level"),
    ("", ""),
    ("Pal Talent Levels", "Per-pal talent skill, one row per star"),
    ("Pal Proficiency", "Per-pal proficiency level curve (exp cost, own_attrs, power)"),
    ("Skill Levels", "Active skill scaling, one row per skill level"),
    ("Mount Ability Curves", "Shared mount ability curves, deduplicated by ability_id (see Mounts.ability_id)"),
    ("Mount Skin Levels", "Per-mount skin level: cost, granted skill levels, attributes, power"),
    ("Star Hero Skill Levels", "Per-hero skill scaling, one row per level"),
    ("Star Hero Star Levels", "Per-hero star-up: cost, attributes, battle skills, develop effect, power"),
    ("Artifact Skin Levels", "Per-artifact skin level: cost, skill levels, attributes, power"),
    ("Back Skin Levels", "Per-back-accessory skin level"),
    ("Avian Affix Levels", "Per-affix level scaling"),
    ("", ""),
    ("attr / own_attrs / value_plus", "Rendered as key(id)=value. num_type 2 attributes are /10000."),
    ("skillCoefficient", "Raw XOR-encoded config values; use desc_parm for the human-readable numbers."),
])

rows = []
for p in pals["pals"]:
    for tal in p.get("talent_progression") or []:
        for lv in tal.get("levels") or []:
            rows.append([p["name"], p["pet_id"], tal.get("skill_id"), tal.get("skill_name"),
                         tal.get("max_star"), lv.get("star"), lv.get("name"),
                         cell("; ".join(lv.get("effect_desc") or [])), cell(lv.get("desc_parm1")),
                         idpairs(lv.get("effect")), cell("; ".join(lv.get("power_desc") or [])),
                         cell(lv.get("desc_parm2")), idpairs(lv.get("power"))])
sheet(wb, "Pal Talent Levels",
      ["pal", "pet_id", "skill_id", "skill_name", "max_star", "star", "level_name", "effect_desc",
       "desc_parm1", "effect", "power_desc", "desc_parm2", "power"], rows,
      widths={"effect_desc": 60, "power_desc": 60, "pal": 22})

rows = []
for p in pals["pals"]:
    for lv in p.get("proficiency_scaling") or []:
        rows.append([p["name"], p["pet_id"], lv.get("level"), lv.get("extra_star"), lv.get("exp"),
                     lv.get("addexp"), pairs(lv.get("own_attrs")), lv.get("power")])
sheet(wb, "Pal Proficiency",
      ["pal", "pet_id", "level", "extra_star", "exp", "addexp", "own_attrs", "power"], rows,
      widths={"own_attrs": 50, "pal": 22})

rows = []
for s in skills["skills"]:
    for lv in s.get("level_scaling") or []:
        rows.append([s["name"], s["skill_id"], lv.get("level"), cell(lv.get("skillCoefficient")),
                     cell(lv.get("desc_parm")), lv.get("power")])
sheet(wb, "Skill Levels", ["skill", "skill_id", "level", "skillCoefficient", "desc_parm", "power"], rows,
      widths={"skill": 24})

rows, seen = [], set()
for m in mounts["mounts"] + mounts["config_only_mounts"]:
    ap = m.get("ability_progression") or {}
    aid = ap.get("ability_id")
    if aid in seen or not ap:
        continue
    seen.add(aid)
    for lv in ap.get("levels") or []:
        rows.append([aid, ap.get("_note", ""), lv.get("level"), pairs(lv.get("value_plus")), lv.get("power")])
sheet(wb, "Mount Ability Curves", ["ability_id", "note", "level", "value_plus", "power"], rows,
      widths={"note": 56, "value_plus": 40})

rows = []
for m in mounts["mounts"]:
    for lv in m.get("skin_progression") or []:
        rows.append([m["name"], m["mount_id"], lv.get("skin_level"), idpairs(lv.get("expend")),
                     idpairs(lv.get("skin_skill")), pairs(lv.get("attr")), lv.get("power")])
sheet(wb, "Mount Skin Levels", ["mount", "mount_id", "skin_level", "expend", "skin_skill", "attr", "power"],
      rows, widths={"mount": 24, "attr": 40})

rows = []
for h in heroes["star_heroes"] + heroes["config_only_angels"]:
    for sk in h.get("skills") or []:
        for lv in sk.get("levels") or []:
            rows.append([h["name"], h.get("angel_id"), sk.get("skill_id"), sk.get("skill_name"),
                         sk.get("type"), sk.get("releaseTime"), sk.get("releaseInterval"),
                         cell(sk.get("rangeType")), lv.get("level"), cell(lv.get("description")),
                         cell(lv.get("desc_parm")), cell(lv.get("skillCoefficient")),
                         pairs(lv.get("ownEffect")), cell(lv.get("attrType")), idpairs(lv.get("expend")),
                         lv.get("power")])
sheet(wb, "Star Hero Skill Levels",
      ["hero", "angel_id", "skill_id", "skill_name", "skill_type", "releaseTime", "releaseInterval",
       "rangeType", "level", "description", "desc_parm", "skillCoefficient", "ownEffect", "attrType",
       "expend", "power"], rows, widths={"description": 70, "hero": 24})

rows = []
for h in heroes["star_heroes"] + heroes["config_only_angels"]:
    for st in h.get("star_progression") or []:
        bs2 = st.get("battle_skill2")
        bs2s = cell(bs2) if not isinstance(bs2, dict) else \
            f"{bs2.get('skill_name')} ({bs2.get('angel_skill_id')}): {bs2.get('description')} | parm={bs2.get('desc_parm')}"
        rows.append([h["name"], h.get("angel_id"), st.get("star"), idpairs(st.get("expend")),
                     pairs(st.get("attr")), st.get("skill1_type"), idpairs(st.get("battle_skill1")),
                     st.get("battle_skill1_cost"), st.get("skill2_type"), cell(bs2s),
                     st.get("battle_skill2_cost"), cell(st.get("develop_effect")),
                     cell(st.get("develop_desc")), cell(st.get("develop_desc_num")),
                     st.get("develop_cost"), st.get("power")])
sheet(wb, "Star Hero Star Levels",
      ["hero", "angel_id", "star", "expend", "attr", "skill1_type", "battle_skill1", "battle_skill1_cost",
       "skill2_type", "battle_skill2", "battle_skill2_cost", "develop_effect", "develop_desc",
       "develop_desc_num", "develop_cost", "power"], rows,
      widths={"battle_skill2": 70, "attr": 44, "hero": 24})

rows = []
for a in artifacts["artifacts"]:
    for lv in a.get("skin_progression") or []:
        rows.append([a["name"], a["artifact_id"], lv.get("skin_level"), idpairs(lv.get("expend")),
                     idpairs(lv.get("skin_skill")), pairs(lv.get("attr")), lv.get("power")])
sheet(wb, "Artifact Skin Levels",
      ["artifact", "artifact_id", "skin_level", "expend", "skin_skill", "attr", "power"], rows,
      widths={"artifact": 26, "attr": 40})

rows = []
for b in backs["back_accessories"]:
    for lv in b.get("skin_progression") or []:
        rows.append([b["name"], b["back_id"], lv.get("skin_level"), idpairs(lv.get("expend")),
                     idpairs(lv.get("skin_skill")), pairs(lv.get("attr")), lv.get("power")])
sheet(wb, "Back Skin Levels",
      ["back_accessory", "back_id", "skin_level", "expend", "skin_skill", "attr", "power"], rows,
      widths={"back_accessory": 26, "attr": 40})

rows = []
for a in affixes["avian_affixes"]:
    for lv in a.get("level_scaling") or []:
        rows.append([a["id"], a["name"], a.get("entry_id"), a.get("rarity"), lv.get("level"),
                     idpairs(lv.get("passive_skill")), cell(lv.get("desc_parm")), lv.get("power")])
sheet(wb, "Avian Affix Levels",
      ["affix_id", "name", "entry_id", "rarity", "level", "passive_skill", "desc_parm", "power"], rows,
      widths={"name": 26})

wb.save(os.path.join(OUT, "LOM_Database_Progressions.xlsx"))
print("wrote LOM_Database_Progressions.xlsx")

# ================================================================= 3. BOSSES
wb = Workbook()
wb.remove(wb.active)
meta = bosses["_meta"]
lines = [("Legend of Mushroom — PvE Boss Database", ""), ("Generated", TODAY)]
for k, v in meta.items():
    if isinstance(v, dict):
        lines.append((k, ""))
        for kk, vv in v.items():
            lines.append(("   " + kk, cell(vv)))
    else:
        lines.append((k, cell(v)))
lines.append(("", ""))
for k, v in bosses.items():
    if k == "_meta":
        continue
    desc = v.get("_description", "") if isinstance(v, dict) else ""
    lines.append((k, cell(desc)))
readme(wb, "README", lines)


def stat_columns(records, statkey="stats"):
    cols = []
    for r in records:
        for k in (r.get(statkey) or {}):
            if k not in cols:
                cols.append(k)
    return cols


def boss_sheet(title, records, base_cols, note=None):
    stats = stat_columns(records)
    rows = []
    for r in records:
        st = r.get("stats") or {}
        rows.append([cell(r.get(c)) for c in base_cols] + [cell(st.get(c)) for c in stats])
    sheet(wb, title, base_cols + stats, rows, note=note)


# story chapters: aggregate ranges, not per-unit
sc = bosses["story_chapters"]
cols = [k for k in sc[0] if k != "sample_boss"]
for r in sc:
    for k in r:
        if k not in cols and k != "sample_boss":
            cols.append(k)
scols = []
for r in sc:
    for k in (r.get("sample_boss") or {}):
        if k not in scols:
            scols.append(k)
rows = []
for r in sc:
    row = []
    for c in cols:
        v = r.get(c)
        if isinstance(v, dict) and set(v) == {"min", "max"}:
            row.append(f"{v['min']} - {v['max']}")
        else:
            row.append(cell(v))
    sb = r.get("sample_boss") or {}
    row += [cell(sb.get(c)) for c in scols]
    rows.append(row)
sheet(wb, "Story Chapters", cols + ["sample_" + c for c in scols], rows,
      note="Per-chapter aggregate: hp/att are min - max across the chapter's boss units; sample_* columns are one representative boss unit.")

boss_sheet("World Boss", bosses["world_boss"]["levels"],
           ["level", "index", "boss_id", "model", "time_limit"],
           note=bosses["world_boss"]["_description"])
boss_sheet("Palu World Boss", bosses["palu_world_boss"]["levels"],
           ["level", "index", "boss_id", "model", "time_limit"],
           note=bosses["palu_world_boss"]["_description"])
boss_sheet("Dark Trial", bosses["dark_trial"]["stages"],
           ["stage_id", "difficulty", "level", "part_type", "model", "primary_boss_id",
            "all_monster_refresh", "time_limit"],
           note=bosses["dark_trial"]["_description"])
boss_sheet("Reversion War Bosses", bosses["reversion_war_boss"]["bosses"],
           [k for k in bosses["reversion_war_boss"]["bosses"][0] if k != "stats"],
           note=bosses["reversion_war_boss"]["_description"])
boss_sheet("League GvE Monsters", bosses["league_gve_chapter_monster"]["monsters"],
           [k for k in bosses["league_gve_chapter_monster"]["monsters"][0] if k != "stats"],
           note=bosses["league_gve_chapter_monster"]["_description"])
boss_sheet("Monster PvE Chapter", bosses["monster_pve_chapter"]["bosses"],
           [k for k in bosses["monster_pve_chapter"]["bosses"][0] if k != "stats"],
           note=bosses["monster_pve_chapter"]["_description"])

rows = []
for b in bosses["activity_bosses"]["bosses"]:
    for lv in b.get("hp_scaling") or [{}]:
        rows.append([b.get("id"), b.get("activity_type"), b.get("activity_type_id"), b.get("boss_id"),
                     b.get("model"), b.get("boss_attack_value"), lv.get("level"), lv.get("hp")])
sheet(wb, "Activity Bosses",
      ["id", "activity_type", "activity_type_id", "boss_id", "model", "boss_attack_value", "level", "hp"],
      rows, note=bosses["activity_bosses"]["_description"])

ae = bosses["alchemy_enemy"]["bosses"]
cols = []
for r in ae:
    for k in r:
        if k not in cols:
            cols.append(k)
sheet(wb, "Alchemy Enemy", cols, [[cell(r.get(c)) for c in cols] for r in ae],
      note=bosses["alchemy_enemy"]["_description"])

wb.save(os.path.join(OUT, "LOM_Database_Bosses.xlsx"))
print("wrote LOM_Database_Bosses.xlsx")
